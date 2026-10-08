from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("프로젝트파일/인터뷰자료/지자체 관광부서 영상 인터뷰 대본과 서비스 사용 의향 확인서.docx")
# macOS 에 기본으로 있는 한글 글꼴 — 설치 안 된 글꼴이면 PDF 변환에서 한글이 빠진다.
FONT = "AppleGothic"
NAVY = "18324B"
BLUE = "EAF2F8"
PALE = "F7F9FB"
BORDER = "D9D9D9"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths=None):
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = grid.rows[0].cells[i]
        cell.text = ""
        shade(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_font(run, 9.5, True, "FFFFFF")
    for ridx, row in enumerate(rows):
        cells = grid.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            if ridx % 2 == 1:
                shade(cell, PALE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(value))
            set_font(run, 9.2)
    if widths:
        for row in grid.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return grid


def paragraph(doc, text="", bold_lead=None, indent=0, after=7, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.38
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        set_font(run, size, True)
        run = p.add_run(text[len(bold_lead):])
        set_font(run, size)
    else:
        run = p.add_run(text)
        set_font(run, size)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, 15 if level == 1 else 12, True, "000000")
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.18 * level)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, 10.3)
    return p


def script_question(doc, no, question, guide, sample):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"Q{no}. {question}")
    set_font(run, 11.2, True, NAVY)
    paragraph(doc, f"진행자 안내: {guide}", bold_lead="진행자 안내: ", indent=0.12, after=3, size=9.7)
    paragraph(doc, f"답변 예시: {sample}", bold_lead="답변 예시: ", indent=0.12, after=5, size=10.2)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("왜곡 지역 설화 기반 방탈출 서비스 인터뷰 자료")
    set_font(footer_run, 8.5, False, "666666")

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(6)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("지자체 관광부서 영상 인터뷰 대본과 서비스 사용 의향 확인서")
    set_font(title_run, 21, True, "000000")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    sub_run = sub.add_run("나주와 목포 관광부서 담당 주무관 및 팀장용")
    set_font(sub_run, 11.5, False, "555555")

    table(doc, ["문서 목적", "사용 방식"], [[
        "지역 설화 기반 방탈출 게임 제작·운영 서비스의 고도화 방향을 소개하고, 영상 인터뷰와 조건부 사용 의향 확인에 활용합니다.",
        "기관별로 지역명과 참여자 이름만 바꿔 사용합니다. 답변 예시는 사실과 다를 때 반드시 참여자 본인의 말로 수정합니다.",
    ]], [3.3, 3.3])

    paragraph(doc, "이 문서는 MVP의 현재 기능을 설명하는 자료가 아닙니다. 지역 설화와 장소 자료를 바탕으로 게임 공간을 제작하고, 최소한의 담당자 검수로 콘텐츠를 운영하며, 관리자가 성과를 비교·판단할 수 있는 서비스의 확장 방향을 설명합니다.")
    paragraph(doc, "영상은 서비스의 필요성과 도입 조건을 듣기 위한 인터뷰입니다. 답변자에게 효과를 단정하거나 구매·도입을 약속하도록 요청하지 않습니다.")

    heading(doc, "1 촬영 전 공통 안내")
    table(doc, ["항목", "권장 기준"], [
        ("영상 길이", "담당 주무관 4~6분, 팀장 5~7분. 질문은 모두 답하지 않아도 됩니다."),
        ("발화 원칙", "개인 또는 기관의 공식 확정 입장으로 오해될 수 있는 표현은 피하고, 실제 판단과 조건을 말합니다."),
        ("소개 범위", "MVP 화면이나 미완성 기능 대신 고도화 서비스의 제작·검수·운영·분석 흐름을 소개합니다."),
        ("확인할 점", "설화 사실성, 지역 맥락, 저작권·초상권, 안전성, 예산·조달, 개인정보, 성과 측정 조건을 함께 확인합니다."),
        ("촬영 방법", "질문을 읽은 뒤 3~5초 쉬고 답합니다. 답변이 길면 문장을 끊어 여러 번 촬영해도 됩니다."),
    ], [1.25, 5.35])

    heading(doc, "2 서비스 고도화 방향 소개문")
    paragraph(doc, "아래 문안은 진행자가 영상 시작 전 약 60초 동안 읽습니다.")
    paragraph(doc, "저희는 지역에 전해지는 설화와 실제 장소를 단순한 안내문이 아니라, 플레이어가 탐색하고 단서를 연결하는 방탈출 경험으로 바꾸는 서비스를 기획하고 있습니다. 지역 담당자는 검토된 설화 자료, 연결할 장소, 이용 가능한 이미지·모델·음원, 지역에서 피해야 할 표현과 안전 기준을 등록합니다. 시스템은 이 범위 안에서 게임 공간과 단서 흐름의 초안을 제안하고, 담당자는 사실성·문화적 적절성·안전성·지역 이미지 관점에서 필요한 부분만 검수합니다.")
    paragraph(doc, "운영 이후에는 관리자 대시보드에서 설화·지역·게임 공간별 플레이 시작, 퍼즐 해결, 힌트 사용, 이탈, 방문 의사와 같은 흐름을 확인합니다. 향후 동의를 받은 파일럿에서는 보상 없음, 식사 할인권, 지역화폐 금액, 숙박 할인권처럼 보상 방식이 다른 경우의 플레이·방문 지표 차이도 비교할 수 있도록 설계합니다. 이 비교는 실제 현장 적용과 적절한 조사 설계를 전제로 하며, 방문 전환을 자동으로 보장한다는 뜻은 아닙니다.")

    heading(doc, "3 우리가 확인하려는 문제와 가설")
    paragraph(doc, "아래는 팀이 정리한 페인포인트와 페르소나를 바탕으로 만든 가설입니다. 인터뷰에서는 ‘맞는지 확인해 주십시오’라고 묻는 방식으로 사용합니다.")
    table(doc, ["문제 가설", "관련 이용자 관점", "고도화 서비스의 대응"], [
        ("설화 자료가 흩어져 있고, 관광 콘텐츠로 재해석하는 데 담당자 시간과 제작 비용이 많이 든다.", "지역 담당자와 콘텐츠 기획자는 사실성·지역성·일관성을 함께 지켜야 한다.", "검토한 설화·장소·자산·금지 표현을 기준으로 제작 범위를 제한하고, 담당자는 핵심 검수에 집중한다."),
        ("설명이 앞서는 관광 콘텐츠는 게임 이용자에게 쉽게 정보 전달처럼 느껴질 수 있다.", "스토리·미스터리·시각적 첫인상으로 유입되는 이용자는 지역 정보를 단서와 사건의 일부로 만나길 원한다.", "설화·장소 정보를 설명문이 아니라 탐색 동기, 퍼즐 규칙, 사건 해결의 근거로 배치한다."),
        ("지역별·설화별 콘텐츠가 실제로 얼마나 플레이되고 방문 의사로 이어지는지 비교하기 어렵다.", "지자체 담당자와 팀장은 재미뿐 아니라 운영 근거, 예산 효율, 재사용 가능성을 확인해야 한다.", "관리자 대시보드에서 진행·이탈·힌트·완주·방문 의사를 지역·사례·보상 조건별로 비교한다."),
        ("보상은 방문 동기를 줄 수 있으나, 어떤 수준과 방식이 적절한지 근거 없이 정하기 어렵다.", "여행 경험형 이용자는 현장 확인과 실질적 혜택을 함께 기대할 수 있다.", "승인된 파일럿에서 보상 조건을 분리해 결과를 비교하고, 과도한 보상이나 지역 특성에 맞지 않는 보상은 줄인다."),
    ], [2.05, 2.05, 2.5])

    heading(doc, "4 담당 주무관 영상 인터뷰 대본")
    paragraph(doc, "아래 ‘답변 예시’는 읽기용 초안입니다. 실제 업무 경험과 다르면 삭제하거나 바꾸어야 합니다.")
    script_question(doc, 1, "지역 설화·관광 콘텐츠를 운영하거나 새로 만들 때 가장 어려운 점은 무엇입니까?", "현재 업무에서 겪는 자료 정리, 사실 확인, 콘텐츠화, 지역 협력, 홍보의 어려움 중 실제 사례를 한두 가지 말해 달라고 요청합니다.", "지역 이야기는 자료로 남아 있어도 관광객이 흥미를 느낄 콘텐츠로 연결하는 과정이 쉽지 않습니다. 사실과 다른 과장이나 지역의 맥락을 놓치지 않아야 하고, 담당자가 하나하나 확인하려면 시간도 많이 듭니다.")
    script_question(doc, 2, "설화와 실제 장소를 게임 속 사건과 퍼즐로 연결하는 방식이 지역 홍보에 어떤 가능성을 줄 수 있다고 보십니까?", "관광 안내가 아닌 몰입형 경험이라는 점에 대해 의견을 묻습니다.", "단순히 장소 정보를 보여 주는 것보다, 이용자가 이야기를 해결하려고 공간을 탐색하는 과정에서 장소의 의미를 자연스럽게 알게 한다는 점이 의미 있다고 생각합니다. 다만 실제 장소와 지역의 정체성을 정확하게 다루는 검수는 전제되어야 합니다.")
    script_question(doc, 3, "시스템이 게임 공간과 퍼즐 초안을 만들고 담당자는 핵심만 검수하는 과정에 대해 어떻게 생각하십니까?", "자동화의 편의와 함께 반드시 사람이 결정해야 하는 범위를 구분해 달라고 요청합니다.", "초안 제작을 줄여 주는 도구라면 검토해 볼 가치가 있습니다. 하지만 설화의 사실성, 주민에게 민감한 표현, 안전 문제, 지역 이미지처럼 책임이 필요한 부분은 사람이 최종 확인해야 한다고 생각합니다.")
    script_question(doc, 4, "도입을 검토한다면 어떤 검수·권한·안전 장치가 필요합니까?", "자료 출처, 승인 단계, 수정 권한, 개인정보, 예산·조달 절차를 구체적으로 말해 달라고 요청합니다.", "출처와 사용 권한이 명확해야 하고, 공개 전에는 담당 부서가 수정·승인할 수 있어야 합니다. 이용자 데이터는 필요한 범위에서만 익명으로 보고, 현장 보상을 연계한다면 운영 주체와 비용, 부정 이용 대응도 먼저 정리되어야 합니다.")
    script_question(doc, 5, "관리자 대시보드에서 어떤 지표를 우선 확인하고 싶으십니까?", "수치 자체보다 어떤 의사결정에 쓰는지 말해 달라고 요청합니다.", "지역이나 설화별로 어디에서 이용자가 많이 시작하고, 어느 퍼즐에서 막히거나 이탈하는지 보고 싶습니다. 방문 의사와 현장 연계 반응도 함께 확인할 수 있다면 다음 콘텐츠와 홍보 예산의 우선순위를 정하는 데 도움이 될 것 같습니다.")
    script_question(doc, 6, "식사 할인권, 지역화폐, 숙박 할인권 같은 보상 조건을 비교하는 기능은 어떤 기준으로 활용되어야 한다고 보십니까?", "보상이 만능 해법이라는 답이 아니라, 지역에 맞는 적정성·비용·형평성 관점을 묻습니다.", "보상은 지역 상권과 연결될 수 있지만, 보상만을 위해 참여하게 만드는 방식은 조심해야 합니다. 조건별 반응을 투명하게 비교하고, 지역의 예산과 참여 상점, 관광 흐름에 맞는 수준인지 검토한 뒤 제한적으로 시험해 보는 방식이 적절하다고 생각합니다.")
    script_question(doc, 7, "조건이 갖춰진다면 파일럿 또는 서비스 검토에 참여할 의향이 있으십니까?", "확정 도입 약속이 아닌 ‘검토·파일럿 참여 의향’만 묻습니다.", "자료 출처와 검수 권한, 개인정보와 예산 조건이 명확하고 작은 범위에서 결과를 확인할 수 있다면, 지역 설화를 활용한 파일럿을 검토해 볼 의향이 있습니다.")

    heading(doc, "5 팀장 영상 인터뷰 대본")
    paragraph(doc, "팀장 답변은 개별 콘텐츠의 호감보다 정책적 타당성, 확장성, 책임 체계, 성과 판단 기준에 초점을 둡니다.")
    script_question(doc, 1, "지역 설화를 관광 콘텐츠로 활용할 때 지자체 차원에서 가장 중요하게 보는 가치는 무엇입니까?", "지역 정체성, 주민과의 관계, 관광 분산, 지역 소비, 지속 가능성 중 실제 우선순위를 말해 달라고 요청합니다.", "지역 고유의 이야기가 일회성 홍보에 머물지 않고, 방문객이 장소의 의미를 이해하고 지역 안에서 경험과 소비로 이어갈 수 있는 구조가 중요하다고 봅니다. 동시에 지역의 맥락을 왜곡하지 않고 지속적으로 운영할 수 있어야 합니다.")
    script_question(doc, 2, "여러 지역과 설화를 같은 체계에서 제작·운영할 수 있다면 어떤 점을 먼저 검토하시겠습니까?", "확장 가능성과 표준화의 이점, 지역별 차이를 존중하는 조건을 함께 묻습니다.", "공통 제작 절차와 지표가 있으면 사례를 비교하고 운영 노하우를 축적하기 좋습니다. 다만 모든 지역 이야기를 같은 형식으로 평준화하지 않도록 지역별 검수와 표현 기준은 별도로 유지되어야 합니다.")
    script_question(doc, 3, "최소 인력 검수 모델을 도입하려면 사람과 시스템의 역할을 어떻게 나누는 것이 적절합니까?", "자동화가 담당할 반복 작업과 사람이 보유해야 할 결정 권한을 분명히 해 달라고 요청합니다.", "자료를 정리하고 초안을 구성하는 반복 작업은 시스템이 도울 수 있습니다. 하지만 설화의 해석, 공개 가능 여부, 지역의 민감성, 보상 정책과 예산 집행처럼 책임이 따르는 결정은 담당자와 관련 부서가 맡아야 합니다.")
    script_question(doc, 4, "관리자 대시보드가 제공된다면 어떤 비교와 보고가 의사결정에 유용하겠습니까?", "지역·설화·게임 공간·보상 조건별 비교와 기간별 변화, 원자료 접근 원칙을 묻습니다.", "사례별 시작·완주·이탈·힌트 사용·방문 의사를 한 화면에서 보고 싶습니다. 보상 조건에 따른 차이는 단순 순위가 아니라 비용과 표본 규모, 운영 기간까지 함께 확인할 수 있어야 다음 예산이나 협력 사업을 판단하는 근거가 됩니다.")
    script_question(doc, 5, "서비스가 실제 방문과 지역 소비에 기여했는지 판단하려면 어떤 검증이 필요합니까?", "‘보장’이 아니라 필요한 검증 설계를 묻습니다.", "온라인 플레이 수만으로 실제 방문 효과를 단정할 수는 없습니다. 사전·사후 설문, 방문 의사, 동의 기반 현장 인증, 협력 상점의 사용 데이터처럼 가능한 지표를 단계적으로 모아야 하고, 비교 대상과 기간도 정해 두어야 합니다.")
    script_question(doc, 6, "나주 또는 목포에서 우선 시험한다면 어떤 범위와 조건이 적절하다고 보십니까?", "작은 단위의 설화·장소·협력처, 담당 부서, 검수 기간, 성과 기준을 제안하도록 요청합니다.", "처음부터 넓은 권역을 대상으로 하기보다 설화와 실제 장소의 연결이 분명한 한 사례에서 시작하는 편이 좋겠습니다. 담당 부서의 검수 범위와 협력처, 보상 예산, 개인정보 처리, 성공·중단 기준을 먼저 합의한 뒤 결과를 보고 확대 여부를 결정해야 합니다.")
    script_question(doc, 7, "이 서비스의 고도화 방향을 검토하거나 파일럿 논의에 참여할 의향이 있으십니까?", "조건부 의향과 필요한 사전 자료를 확인합니다.", "지역성 보장, 검수 체계, 예산과 개인정보 기준, 성과 측정 방법이 정리된다면 파일럿 가능성을 검토해 볼 수 있습니다. 도입 여부는 실제 제안서와 내부 검토 절차를 거쳐 판단하겠습니다.")

    heading(doc, "6 서비스 사용 의향 확인서")
    paragraph(doc, "이 확인서는 영상 인터뷰와 함께 사용할 수 있는 비구속적 의향 확인 양식입니다. 계약, 예산 집행, 사업 참여 확정 또는 기관의 공식 입장을 대체하지 않습니다.")
    table(doc, ["확인 항목", "기입 또는 선택"], [
        ("기관명", "____________________________________________"),
        ("부서명", "____________________________________________"),
        ("성명 및 직위", "____________________________________________"),
        ("대상 지역", "□ 나주  □ 목포  □ 기타 ________________________"),
        ("검토 의향", "□ 조건 충족 시 파일럿 검토 의향 있음  □ 추가 자료 검토 후 판단  □ 현재 검토 의향 없음"),
        ("관심 기능", "□ 설화·장소 기반 게임 제작  □ 담당자 검수 흐름  □ 관리자 대시보드  □ 보상 조건 비교  □ 현장 연계"),
        ("도입 전 필요 조건", "□ 출처·권리 확인  □ 검수 권한  □ 개인정보 기준  □ 예산·조달 검토  □ 성과 측정 설계  □ 기타 __________"),
        ("검토 의견", "________________________________________________________________________________\n________________________________________________________________________________"),
        ("작성일", "20____년 ____월 ____일"),
        ("확인자", "성명 ____________________  서명 또는 날인 ____________________"),
    ], [1.65, 4.95])

    heading(doc, "7 영상 촬영 협조 및 활용 범위 확인")
    paragraph(doc, "아래 항목은 소속 기관의 내부 규정에 맞게 조정해 사용합니다. 별도의 공식 동의서가 필요한 경우에는 기관 양식을 우선합니다.")
    table(doc, ["항목", "확인 내용"], [
        ("촬영 목적", "프로젝트 연구·발표·서비스 검토 자료를 위한 인터뷰 영상"),
        ("활용 범위", "□ 내부 검토  □ 수업·경진대회 발표  □ 프로젝트 소개 자료  □ 기타 ____________________"),
        ("편집 동의", "□ 발언 의미를 바꾸지 않는 범위의 길이 편집에 동의  □ 게시 전 최종 확인 필요"),
        ("공개 표기", "□ 소속·직위 공개 가능  □ 성명 비공개  □ 소속 비공개"),
        ("확인", "위 활용 범위를 확인했습니다.  성명 ____________________  날짜 ____________________"),
    ], [1.45, 5.15])

    heading(doc, "8 촬영 후 정리 체크리스트")
    for item in [
        "답변자가 사실과 다른 예시 문장을 그대로 읽지 않았는지 확인한다.",
        "‘방문 전환 보장’, ‘도입 확정’, ‘공식 승인’처럼 확인되지 않은 표현을 삭제한다.",
        "기관명·직위·활용 범위 공개 여부를 다시 확인한다.",
        "사용 의향 확인서는 서명 여부와 관계없이 계약이나 예산 확정으로 해석하지 않는다.",
        "영상 편집본은 발언의 앞뒤 맥락이 바뀌지 않도록 확인한다.",
    ]:
        bullet(doc, "□ " + item)

    doc.core_properties.title = "지자체 관광부서 영상 인터뷰 대본과 서비스 사용 의향 확인서"
    doc.core_properties.subject = "나주와 목포 관광부서용"
    doc.core_properties.author = "Latent Space"
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
